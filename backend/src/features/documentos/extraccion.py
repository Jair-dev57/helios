from __future__ import annotations
from pathlib import Path

LIMITE_PALABRAS = 6000
BASE_DIR = Path(__file__).resolve().parent.parent.parent.parent  # backend/


def _ruta_local(ruta_relativa: str) -> Path:
    """Convierte '/uploads/documentos/xxx.pdf' en la ruta real en disco."""
    return BASE_DIR / ruta_relativa.lstrip("/")


def _limpiar_texto(texto: str) -> str:
    """Elimina caracteres de control no imprimibles que a veces deja pypdf (viñetas, etc.)."""
    return "".join(c for c in texto if c.isprintable() or c in "\n\t")


def _extraer_pdf(ruta: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(ruta))
    texto = "\n".join(page.extract_text() or "" for page in reader.pages)
    return _limpiar_texto(texto)

def _extraer_docx(ruta: Path) -> str:
    from docx import Document
    doc = Document(str(ruta))
    return "\n".join(p.text for p in doc.paragraphs)


def _extraer_xlsx(ruta: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(ruta), data_only=True, read_only=True)
    partes = []
    for hoja in wb.worksheets:
        for fila in hoja.iter_rows(values_only=True):
            valores = [str(c) for c in fila if c is not None]
            if valores:
                partes.append(" | ".join(valores))
    return "\n".join(partes)


def _extraer_texto_plano(ruta: Path) -> str:
    return ruta.read_text(encoding="utf-8", errors="ignore")


EXTRACTORES = {
    "pdf": _extraer_pdf,
    "docx": _extraer_docx,
    "doc": _extraer_docx,
    "xlsx": _extraer_xlsx,
    "xls": _extraer_xlsx,
    "csv": _extraer_texto_plano,
    "txt": _extraer_texto_plano,
}


def extraer_contenido(ruta_relativa: str, tipo: str, limite_palabras: int = LIMITE_PALABRAS) -> str | None:
    """Extrae el texto de un archivo según su tipo. Retorna None si el tipo no es soportado o falla la lectura."""
    extractor = EXTRACTORES.get(tipo.lower().lstrip("."))
    if not extractor:
        return None

    ruta = _ruta_local(ruta_relativa)
    if not ruta.exists():
        return None

    try:
        texto = extractor(ruta)
    except Exception:
        return None

    if not texto or not texto.strip():
        return None

    palabras = texto.split()
    if len(palabras) > limite_palabras:
        texto = " ".join(palabras[:limite_palabras])

    return texto.strip()
